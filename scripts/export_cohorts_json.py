"""
Excel 백테스트 결과 → web/public/data/ JSON 변환 스크립트

생성 파일:
  cohorts_ndx3x_10b.json   — TQQQ A/B/C, 목표 10억
  cohorts_ndx2x_10b.json   — QLD
  cohorts_ndx1x_10b.json   — QQQ
  cohorts_sp500_10b.json   — VOO/SPX

실행: python scripts/export_cohorts_json.py
데이터 업데이트 후 재실행하면 JSON 자동 갱신됨.
"""

import json
import openpyxl
from pathlib import Path

RESULTS_DIR = Path("D:/justkeepbuyingtqqq/results")
OUT_DIR     = Path("D:/justkeepbuyingtqqq/web/public/data")

OUT_DIR.mkdir(parents=True, exist_ok=True)

FILES = [
    ("mcv_3x_compare_A_B_C_10b.xlsx",      "cohorts_ndx3x_10b.json"),
    ("mcv_2x_compare_A_B_C_10b.xlsx",      "cohorts_ndx2x_10b.json"),
    ("mcv_1x_compare_A_B_C_10b.xlsx",      "cohorts_ndx1x_10b.json"),
    ("mcv_spx500_1x_compare_A_B_C_10b.xlsx", "cohorts_sp500_10b.json"),
]

def xlsx_to_json(xlsx_path: Path, json_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    rows = []
    for ws_row in ws.iter_rows(min_row=2, values_only=True):
        raw = dict(zip(headers, ws_row))

        # start_date → "YYYY-MM-DD"
        sd = raw["start_date"]
        start_str = sd.strftime("%Y-%m-%d") if hasattr(sd, "strftime") else str(sd)[:10]

        # end_date를 years로 역산 (end_date 컬럼 없으므로 start + years * 365.25)
        def end_date_str(start, years):
            if years is None:
                return None
            import datetime
            end = start + datetime.timedelta(days=years * 365.25)
            return end.strftime("%Y-%m-%d")

        row = {
            "s":  start_str,
            # A전략
            "sA": raw.get("status_A"),
            "yA": round(raw["years_A"], 4) if raw.get("years_A") is not None else None,
            "eA": end_date_str(sd, raw.get("years_A")),
            "iA": int(raw["invested_A"]) if raw.get("invested_A") is not None else None,
            # B전략
            "sB": raw.get("status_B"),
            "yB": round(raw["years_B"], 4) if raw.get("years_B") is not None else None,
            "eB": end_date_str(sd, raw.get("years_B")),
            "iB": int(raw["invested_B"]) if raw.get("invested_B") is not None else None,
            # C전략
            "sC": raw.get("status_C"),
            "yC": round(raw["years_C"], 4) if raw.get("years_C") is not None else None,
            "eC": end_date_str(sd, raw.get("years_C")),
            "iC": int(raw["invested_C"]) if raw.get("invested_C") is not None else None,
        }
        rows.append(row)

    # 메타 정보 포함
    completed_B = [r for r in rows if r["sB"] == "completed"]
    years_B = [r["yB"] for r in completed_B]

    meta = {
        "params": {"dailyInvest": 200000, "lumpSum": 250000000, "target": 1000000000},
        "total": len(rows),
        "completedB": len(completed_B),
        "avgB":    round(sum(years_B)/len(years_B), 2) if years_B else None,
        "maxB":    round(max(years_B), 2) if years_B else None,
        "source":  xlsx_path.name,
    }

    output = {"meta": meta, "rows": rows}

    json_path.write_text(
        json.dumps(output, ensure_ascii=False, separators=(',', ':')),
        encoding="utf-8"
    )

    size_kb = json_path.stat().st_size / 1024
    print(f"  ✓ {json_path.name}  ({len(rows)}행, {size_kb:.0f} KB)")


if __name__ == "__main__":
    print("JSON 생성 중...\n")
    for xlsx_name, json_name in FILES:
        xlsx_path = RESULTS_DIR / xlsx_name
        json_path = OUT_DIR / json_name
        if not xlsx_path.exists():
            print(f"  ⚠ 파일 없음: {xlsx_path}")
            continue
        xlsx_to_json(xlsx_path, json_path)

    print(f"\n저장 위치: {OUT_DIR}")
    print("웹 접근 경로: /data/<파일명>.json")
