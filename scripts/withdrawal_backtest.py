"""
justkeepbuyingtqqq — 인출(파이어) 단계 백테스트

버퍼 공식: buffer_years = max(0, BUFFER_REF_YEARS - years_to_target)
  BUFFER_REF_YEARS = 13  (NDX3x B 실측 최장 12.24년 올림)

인출 룰:
  트리거: SP500 지수 드로다운 기준
  WITHDRAW  : 월 1% 매도, SP500 -20% → RESCUE_1Y, SP500 -50% → BUFFER_BUY
  RESCUE_1Y : 1년간 일 daily_invest 매수 (외부 자금)
              8억 회복 or 1년 만료 → WITHDRAW
  BUFFER_BUY: 버퍼를 BUFFER_REF_YEARS 년에 걸쳐 분할 매수, 8억 회복 → WITHDRAW
  DCA_ENTRY : 버블 케이스 전용 — 보유 현금을 buffer_years년에 걸쳐 매일 분할 매수
              현금 소진 or 8억 초과 → WITHDRAW

버블 케이스 (years_to_target < 1.5년): 전량매도 후 두 가지 재진입 방식 비교
  Method A: 2년 대기 후 DCA_ENTRY (buffer_years년 분산 매수)
  Method B: 즉시 DCA_ENTRY         (기다림 없이 바로 분산 매수 시작)
"""

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.chart import LineChart, Reference
import time

# ===== 공유 파라미터 =====
DAILY_INVEST           = 200_000
CAP_INVEST             = 250_000_000
LUMP_SUM               = 250_000_000
ACCUM_TARGET           = 1_000_000_000
WITHDRAW_THRESHOLD     = DAILY_INVEST * 4_000   # 8억
MONTHLY_WITHDRAW_RATE  = 0.01
BUFFER_REF_YEARS       = 13
BUBBLE_THRESHOLD_YEARS = 1.5
BUBBLE_WAIT_YEARS      = 2.0                    # Method A 대기 기간
RESCUE_DURATION_DAYS   = 252
DATA_DIR               = "D:/justkeepbuyingtqqq/data/"
OUT_DIR                = "D:/justkeepbuyingtqqq/results/"


# ===== 데이터 로드 =====
def load_ndx(filepath):
    df = pd.read_csv(filepath, parse_dates=["Date"]).sort_values("Date").reset_index(drop=True)
    mask = df["Date"] == "1985-10-01"
    if mask.any():
        i = df.index[mask][0]
        scale = df.loc[i - 1, "Close"] / df.loc[i, "Close"]
        df.loc[i:, "Close"] *= scale
        print(f"  NDX splice 보정: scale={scale:.4f}")
    return df

def load_sp500(filepath):
    return pd.read_csv(filepath, parse_dates=["Date"]).sort_values("Date").reset_index(drop=True)

def make_synthetic(df, leverage):
    ret = df["Close"].pct_change().fillna(0).values
    return 100.0 * np.cumprod(1.0 + ret * leverage)

def align_sp500_to(sp5_df, target_dates):
    series = sp5_df.set_index("Date")["Close"]
    return series.reindex(target_dates, method="ffill").values

def get_monthly_starts(dates_pd):
    seen, starts = set(), []
    for i, d in enumerate(dates_pd):
        k = (d.year, d.month)
        if k not in seen:
            seen.add(k)
            starts.append(i)
    return starts


# ===== 1단계: 적립 백테스트 =====
def run_accumulation(prices, dates_pd, strategy, daily_invest, cap_invest, lump_sum, target):
    n_invest_days = int(cap_invest // daily_invest)
    cohort_starts = get_monthly_starts(dates_pd)
    results = []
    for si in cohort_starts:
        n = len(prices) - si
        start_date = dates_pd[si]
        cum_shares = cum_invest = 0.0
        hit_idx = -1
        hit_val = 0.0
        for j in range(n):
            px = prices[si + j]
            if strategy == "A":
                inv = daily_invest if j < n_invest_days else 0.0
            elif strategy == "B":
                inv = lump_sum if j == 0 else daily_invest
            else:
                inv = daily_invest
            cum_shares += inv / px
            cum_invest += inv
            port_val = cum_shares * px
            if port_val >= target:
                hit_idx = j
                hit_val = port_val
                break
        if hit_idx >= 0:
            end_date = dates_pd[si + hit_idx]
            days = (end_date - start_date).days
            results.append({
                "start_date":      start_date,
                "end_date":        end_date,
                "status":          "completed",
                "years_to_target": days / 365.25,
                "final_value":     hit_val,
                "end_idx":         si + hit_idx,
            })
        else:
            results.append({
                "start_date":      start_date,
                "end_date":        None,
                "status":          "in_progress",
                "years_to_target": None,
                "final_value":     cum_shares * prices[-1],
                "end_idx":         len(prices) - 1,
            })
    return results


# ===== 2단계: 인출 시뮬레이션 =====
def run_withdrawal(
    prices, dates_pd, start_idx, portfolio_value, years_to_target,
    sp500_ref, bubble_mode="A",
    daily_invest=DAILY_INVEST, monthly_rate=MONTHLY_WITHDRAW_RATE, sim_years=40,
):
    """
    bubble_mode: "A" = 2년 대기 후 DCA
                 "B" = 즉시 DCA
    """
    buffer_years  = max(0.0, BUFFER_REF_YEARS - years_to_target)
    buffer_amount = buffer_years * daily_invest * 252

    n_total = len(prices)
    sim_len = min(n_total - start_idx, int(sim_years * 252))
    if sim_len <= 0:
        return None

    bubble_case = years_to_target < BUBBLE_THRESHOLD_YEARS

    if bubble_case:
        dca_days = max(1, int(buffer_years * 252))   # DCA 분산 기간

        if bubble_mode == "A":
            # ── Method A: 2년 대기 후 DCA ──
            wait_days  = int(BUBBLE_WAIT_YEARS * 252)
            re_entry_i = start_idx + wait_days
            if re_entry_i >= n_total:
                return _make_result(
                    [], portfolio_value, portfolio_value, 0.0, 0,
                    "bubble_cash_only", buffer_years, years_to_target,
                    dates_pd, start_idx, sim_len
                )
            new_sim_len = min(n_total - re_entry_i, sim_len - wait_days)
            monthly_data, total_withdrawn, min_val = _simulate_withdraw(
                prices, dates_pd, re_entry_i, new_sim_len,
                shares=0.0, cash_buffer=0.0,
                daily_invest=daily_invest, monthly_rate=monthly_rate,
                sp500_ref=sp500_ref,
                dca_cash=portfolio_value, dca_days=dca_days,
            )
            return _make_result(
                monthly_data, portfolio_value,
                monthly_data[-1]["value"] if monthly_data else 0,
                total_withdrawn, min_val,
                "bubble_A(wait2+dca)", buffer_years, years_to_target,
                dates_pd, start_idx, sim_len
            )

        else:
            # ── Method B: 즉시 DCA ──
            monthly_data, total_withdrawn, min_val = _simulate_withdraw(
                prices, dates_pd, start_idx, sim_len,
                shares=0.0, cash_buffer=0.0,
                daily_invest=daily_invest, monthly_rate=monthly_rate,
                sp500_ref=sp500_ref,
                dca_cash=portfolio_value, dca_days=dca_days,
            )
            return _make_result(
                monthly_data, portfolio_value,
                monthly_data[-1]["value"] if monthly_data else 0,
                total_withdrawn, min_val,
                "bubble_B(immediate_dca)", buffer_years, years_to_target,
                dates_pd, start_idx, sim_len
            )

    # ── 일반 케이스 (버블 아닌 경우 — 두 method 동일) ──
    shares = portfolio_value / prices[start_idx]
    monthly_data, total_withdrawn, min_val = _simulate_withdraw(
        prices, dates_pd, start_idx, sim_len,
        shares=shares, cash_buffer=buffer_amount,
        daily_invest=daily_invest, monthly_rate=monthly_rate,
        sp500_ref=sp500_ref,
    )
    return _make_result(
        monthly_data, portfolio_value,
        monthly_data[-1]["value"] if monthly_data else 0,
        total_withdrawn, min_val,
        "normal", buffer_years, years_to_target,
        dates_pd, start_idx, sim_len
    )


def _simulate_withdraw(
    prices, dates_pd, start_i, sim_len,
    shares, cash_buffer,
    daily_invest, monthly_rate, sp500_ref,
    dca_cash=0.0, dca_days=0,
):
    """
    dca_cash > 0: DCA_ENTRY 상태로 시작 (버블 케이스 전용)
    """
    n_total = len(prices)
    threshold = WITHDRAW_THRESHOLD
    buffer_spread_days = int(BUFFER_REF_YEARS * 252)

    # 초기 상태
    dca_mode = (dca_cash > 0 and dca_days > 0)
    state = "DCA_ENTRY" if dca_mode else "WITHDRAW"

    dca_cash_remain = dca_cash
    dca_daily       = dca_cash / dca_days if dca_days > 0 else 0.0

    # 피크: DCA_ENTRY 시작 시 shares=0이므로 threshold로 초기화 (트리거 오발 방지)
    local_peak   = shares * prices[start_i] if shares > 0 else threshold
    sp500_peak   = sp500_ref[start_i]

    rescue_days_left  = 0
    buffer_days_used  = 0
    cash_buffer_remain = cash_buffer

    total_withdrawn = 0.0
    monthly_snapshots = []
    withdraw_prev_month = -1
    snap_prev_month     = -1

    for j in range(sim_len):
        ci = start_i + j
        if ci >= n_total:
            break
        px      = prices[ci]
        sp_px   = sp500_ref[ci]
        port_val = shares * px

        # 피크 업데이트
        if port_val > local_peak:
            local_peak = port_val
        if sp_px > sp500_peak:
            sp500_peak = sp_px
        sp500_dd = (sp_px / sp500_peak) - 1.0 if sp500_peak > 0 else 0.0

        # ---- 상태 전이 ----
        if state == "DCA_ENTRY":
            if dca_cash_remain <= 0 or port_val >= threshold:
                state = "WITHDRAW"
                local_peak = max(port_val, threshold)
                sp500_peak = sp_px

        elif state == "WITHDRAW":
            if sp500_dd <= -0.50:
                state = "BUFFER_BUY"
                buffer_days_used = 0
            elif sp500_dd <= -0.20:
                state = "RESCUE_1Y"
                rescue_days_left = RESCUE_DURATION_DAYS

        elif state == "RESCUE_1Y":
            if sp500_dd <= -0.50:
                state = "BUFFER_BUY"
                buffer_days_used = 0
            elif port_val > threshold:
                state = "WITHDRAW"
                local_peak = port_val
                sp500_peak = sp_px
            elif rescue_days_left <= 0:
                state = "WITHDRAW"
                local_peak = port_val
                sp500_peak = sp_px

        elif state == "BUFFER_BUY":
            if port_val > threshold:
                state = "WITHDRAW"
                local_peak = port_val
                sp500_peak = sp_px

        # ---- 액션 ----
        cur_month = (dates_pd[ci].year, dates_pd[ci].month)

        if state == "DCA_ENTRY":
            buy = min(dca_daily, dca_cash_remain)
            shares += buy / px
            dca_cash_remain -= buy

        elif state == "WITHDRAW":
            if cur_month != withdraw_prev_month and j > 0:
                withdrawal = shares * px * monthly_rate
                shares = max(0.0, shares - withdrawal / px)
                total_withdrawn += withdrawal
            withdraw_prev_month = cur_month

        elif state == "RESCUE_1Y":
            shares += daily_invest / px
            rescue_days_left -= 1

        elif state == "BUFFER_BUY":
            if cash_buffer_remain > 0 and buffer_days_used < buffer_spread_days:
                remaining_days = buffer_spread_days - buffer_days_used
                daily_buy = cash_buffer_remain / remaining_days
                shares += daily_buy / px
                cash_buffer_remain -= daily_buy
                buffer_days_used += 1

        # ---- 스냅샷 ----
        if cur_month != snap_prev_month:
            monthly_snapshots.append({
                "date":  dates_pd[ci].strftime("%Y-%m"),
                "value": shares * px,
                "state": state,
            })
            snap_prev_month = cur_month

        if shares <= 0:
            break

    min_val = min((s["value"] for s in monthly_snapshots), default=0.0)
    return monthly_snapshots, total_withdrawn, min_val


def _make_result(monthly_data, initial_value, final_value, total_withdrawn,
                 min_value, case_type, buffer_years, years_to_target,
                 dates_pd, start_idx, sim_len):
    survived = final_value > 0
    end_i = min(start_idx + sim_len - 1, len(dates_pd) - 1)
    return {
        "case_type":       case_type,
        "buffer_years":    round(buffer_years, 2),
        "years_to_target": round(years_to_target, 2),
        "initial_value":   int(initial_value),
        "final_value":     int(final_value),
        "total_withdrawn": int(total_withdrawn),
        "min_value":       int(min_value),
        "survived":        survived,
        "sim_end_date":    dates_pd[end_i].strftime("%Y-%m-%d"),
        "monthly_data":    monthly_data,
    }


# ===== 3단계: 파이프라인 =====
def run_full_pipeline(prices, dates_pd, strategy, leverage,
                      daily_invest, cap_invest, lump_sum,
                      accum_target, sim_years, label, out_path,
                      sp500_ref, bubble_mode="A"):
    print(f"\n  [{label}]  bubble_mode={bubble_mode}")
    accum_results = run_accumulation(
        prices, dates_pd, strategy, daily_invest, cap_invest, lump_sum, accum_target
    )
    completed = [r for r in accum_results if r["status"] == "completed"]
    print(f"    완료 코호트: {len(completed)} / {len(accum_results)}")

    rows = []
    for r in completed:
        w = run_withdrawal(
            prices, dates_pd,
            r["end_idx"], r["final_value"], r["years_to_target"],
            sp500_ref=sp500_ref, bubble_mode=bubble_mode,
            daily_invest=daily_invest, monthly_rate=MONTHLY_WITHDRAW_RATE,
            sim_years=sim_years,
        )
        if w:
            rows.append({
                "start_date":        r["start_date"].strftime("%Y-%m-%d"),
                "accum_end_date":    r["end_date"].strftime("%Y-%m-%d"),
                "years_to_target":   round(r["years_to_target"], 3),
                "buffer_years":      w["buffer_years"],
                "case_type":         w["case_type"],
                "initial_portfolio": w["initial_value"],
                "final_value":       w["final_value"],
                "total_withdrawn":   w["total_withdrawn"],
                "min_value":         w["min_value"],
                "survived":          w["survived"],
                "sim_end_date":      w["sim_end_date"],
                "_monthly":          w["monthly_data"],
            })

    df = pd.DataFrame([{k: v for k, v in r.items() if k != "_monthly"} for r in rows])
    sv = df["survived"].sum()
    n  = len(df)
    print(f"    생존율: {sv}/{n} = {sv/n*100:.1f}%")
    print(f"    최솟값(최악): {df['min_value'].min()/1e8:.2f}억  평균: {df['min_value'].mean()/1e8:.2f}억")
    print(f"    최종값 평균: {df['final_value'].mean()/1e8:.0f}억")
    print(f"    인출 총액 평균: {df['total_withdrawn'].mean()/1e8:.1f}억")

    _save_withdrawal_excel(df, rows, out_path, label)
    return df, rows


def _save_withdrawal_excel(df, rows_with_monthly, out_path, label):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "summary"
    cols = ["start_date","accum_end_date","years_to_target","buffer_years",
            "case_type","initial_portfolio","final_value","total_withdrawn",
            "min_value","survived","sim_end_date"]
    ws.append(cols)
    for r in df.itertuples(index=False):
        ws.append([getattr(r, c) for c in cols])

    ws2 = wb.create_sheet("stats")
    n = len(df); sv = df["survived"].sum()
    ws2.append(["지표","값"])
    ws2.append(["총 코호트",n]); ws2.append(["생존",int(sv)])
    ws2.append(["생존율(%)",round(sv/n*100,1)])
    ws2.append(["최솟값 최솟값(억)",round(df["min_value"].min()/1e8,2)])
    ws2.append(["평균 최솟값(억)",round(df["min_value"].mean()/1e8,2)])
    ws2.append(["평균 최종값(억)",round(df["final_value"].mean()/1e8,1)])
    ws2.append(["평균 인출 총액(억)",round(df["total_withdrawn"].mean()/1e8,1)])

    df_sorted_min = df.sort_values("min_value")
    df_sorted_fin = df.sort_values("final_value", ascending=False)
    key = {}
    if len(df_sorted_min): key["worst"]  = df_sorted_min.iloc[0]["start_date"]
    if len(df) > 1:        key["median"] = df.sort_values("years_to_target").iloc[len(df)//2]["start_date"]
    if len(df_sorted_fin): key["best"]   = df_sorted_fin.iloc[0]["start_date"]

    ws3 = wb.create_sheet("monthly_key_cohorts")
    ws3.append(["date","worst(억)","median(억)","best(억)"])
    monthly_by_start = {r["start_date"]: r["_monthly"] for r in rows_with_monthly}
    worst_dates = [m["date"] for m in monthly_by_start.get(key.get("worst"), [])]
    worst_vals  = [m["value"]/1e8 for m in monthly_by_start.get(key.get("worst"), [])]
    med_vals    = {m["date"]: m["value"]/1e8 for m in monthly_by_start.get(key.get("median"), [])}
    best_vals   = {m["date"]: m["value"]/1e8 for m in monthly_by_start.get(key.get("best"), [])}
    for i, d in enumerate(worst_dates):
        ws3.append([d, worst_vals[i] if i < len(worst_vals) else None,
                    med_vals.get(d), best_vals.get(d)])

    if len(worst_dates) > 1:
        chart = LineChart()
        chart.title = f"{label} 인출 포트폴리오 (억원)"
        chart.style = 10; chart.y_axis.title = "포트폴리오 (억원)"
        nrows = len(worst_dates)
        for ci2, sn in [(2,"최악"),(3,"중간"),(4,"최선")]:
            data = Reference(ws3, min_col=ci2, max_col=ci2, min_row=1, max_row=nrows+1)
            chart.series.append(openpyxl.chart.Series(data, title=sn))
        chart.width = 24; chart.height = 14
        ws3.add_chart(chart, "F2")

    wb.save(out_path)
    print(f"    → {out_path.split('/')[-1]}")


# ===== 메인 =====
def main():
    t0 = time.time()
    print("=" * 60)
    print("justkeepbuyingtqqq — 인출 백테스트  버블 재진입 비교")
    print(f"  BUFFER_REF={BUFFER_REF_YEARS}년  트리거=SP500  threshold={BUBBLE_THRESHOLD_YEARS}년")
    print("=" * 60)

    print("\n[데이터 로드]")
    ndx_df = load_ndx(DATA_DIR + "ndx_1971_now.csv")
    sp5_df = load_sp500(DATA_DIR + "sp500_1927_now.csv")
    ndx_dates = pd.DatetimeIndex(ndx_df["Date"])
    ndx3x = make_synthetic(ndx_df, 3)
    sp500_for_ndx = align_sp500_to(sp5_df, ndx_dates)
    print(f"  NDX: {len(ndx_df)}행  SP500: {len(sp5_df)}행  NaN: {np.isnan(sp500_for_ndx).sum()}")

    common = dict(
        prices=ndx3x, dates_pd=ndx_dates, strategy="B", leverage="3x",
        daily_invest=DAILY_INVEST, cap_invest=CAP_INVEST, lump_sum=LUMP_SUM,
        accum_target=ACCUM_TARGET, sim_years=40, sp500_ref=sp500_for_ndx,
    )

    print("\n── Method A: 2년 대기 후 buffer_years DCA ──")
    df_A, rows_A = run_full_pipeline(
        **common, label="NDX3x_B_bubbleA",
        out_path=OUT_DIR + "mcv_withdrawal_ndx3x_B_bubbleA.xlsx",
        bubble_mode="A",
    )

    print("\n── Method B: 즉시 buffer_years DCA ──")
    df_B, rows_B = run_full_pipeline(
        **common, label="NDX3x_B_bubbleB",
        out_path=OUT_DIR + "mcv_withdrawal_ndx3x_B_bubbleB.xlsx",
        bubble_mode="B",
    )

    # ── 1996-10 코호트 비교 ──
    print("\n" + "=" * 60)
    print("1996-10 코호트 (닷컴버블 최악 시나리오) 비교")
    print("=" * 60)
    for label, df in [("Method A", df_A), ("Method B", df_B)]:
        row = df[df["start_date"] == "1996-10-01"]
        if not row.empty:
            r = row.iloc[0]
            print(f"  {label}: min={r.min_value/1e8:.2f}억  final={r.final_value/1e8:.0f}억  "
                  f"total_out={r.total_withdrawn/1e8:.1f}억  case={r.case_type}")

    # ── 버블 케이스 전체 비교 ──
    print("\n버블 케이스 (years_to_target < 1.5) 전체:")
    for label, df in [("Method A", df_A), ("Method B", df_B)]:
        b = df[df["case_type"].str.startswith("bubble")]
        sv = b["survived"].sum()
        print(f"  {label}: {sv}/{len(b)} 생존  "
              f"min최솟값={b.min_value.min()/1e8:.2f}억  "
              f"min평균={b.min_value.mean()/1e8:.2f}억  "
              f"final평균={b.final_value.mean()/1e8:.0f}억")

    print(f"\n총 소요: {time.time()-t0:.1f}초")


if __name__ == "__main__":
    main()
